from __future__ import annotations

import re
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
)


def _text(value: object, fallback: str = "A definir mediante validação") -> str:
    cleaned = str(value or "").strip()
    if not cleaned:
        return fallback
    cleaned = re.sub(
        r"\bespens_clean\.xlsx\b",
        "Ações Educativas ESPEN",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"\b(?:MCN\s*2026|Matriz Curricular Nacional\s*-?\s*2026)\b",
        "Matriz Curricular Nacional - 2026",
        cleaned,
        flags=re.IGNORECASE,
    )
    return cleaned


def _items(values: object) -> list[str]:
    if not isinstance(values, list):
        return ["A definir mediante validação"]
    cleaned = [_text(value, "") for value in values]
    result = [value for value in cleaned if value]
    return result or ["A definir mediante validação"]


def _set_docx_font(run, name: str = "Arial") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def _add_docx_labeled_paragraph(document: Document, label: str, value: object) -> None:
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    _set_docx_font(label_run)
    value_run = paragraph.add_run(_text(value))
    _set_docx_font(value_run)


def _add_docx_list(document: Document, values: object) -> None:
    for value in _items(values):
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(value)
        _set_docx_font(run)


def _configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    title_borders = styles["Title"]._element.pPr.find(qn("w:pBdr"))
    if title_borders is not None:
        styles["Title"]._element.pPr.remove(title_borders)


def build_teaching_plan_docx(plan: dict) -> bytes:
    document = Document()
    _configure_docx(document)

    title = document.add_paragraph(style="Title")
    title.add_run(f"Plano de Ensino {_text(plan.get('title'))}")

    document.add_heading("Identificação e justificativa", level=1)
    _add_docx_labeled_paragraph(document, "Público alvo", plan.get("target_audience"))
    _add_docx_labeled_paragraph(document, "Modalidade", plan.get("modality"))
    _add_docx_labeled_paragraph(document, "Carga horária total", plan.get("total_workload"))
    _add_docx_labeled_paragraph(document, "Justificativa", plan.get("rationale"))
    _add_docx_labeled_paragraph(document, "Desempenho esperado", plan.get("expected_performance"))

    document.add_heading("Objetivos", level=1)
    _add_docx_labeled_paragraph(document, "Objetivo geral", plan.get("general_objective"))
    document.add_heading("Objetivos específicos", level=2)
    _add_docx_list(document, plan.get("specific_objectives"))

    document.add_heading("Trilha de aprendizagem", level=1)
    steps = plan.get("learning_path") if isinstance(plan.get("learning_path"), list) else []
    if not steps:
        document.add_paragraph("A definir mediante validação")
    for index, step in enumerate(steps, start=1):
        sequence = step.get("sequence", index)
        document.add_heading(
            f"Etapa {sequence} {_text(step.get('title'))}",
            level=2,
        )
        _add_docx_labeled_paragraph(document, "Objetivo", step.get("learning_objective"))
        for label, key in (
            ("Conteúdos", "contents"),
            ("Metodologia", "methodology"),
            ("Avaliação", "assessment"),
            ("Alinhamento com a MCN", "mcn_alignment"),
        ):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(label)
            run.bold = True
            _set_docx_font(run)
            _add_docx_list(document, step.get(key))
        _add_docx_labeled_paragraph(document, "Carga horária", step.get("workload"))

    for heading, key in (
        ("Estratégias de ensino", "teaching_methods"),
        ("Recursos necessários", "resources"),
        ("Avaliação da aprendizagem", "learning_assessment"),
        ("Avaliação de transferência e impacto", "transfer_impact_assessment"),
        ("Critérios de certificação", "certification_criteria"),
        ("Referências", "references"),
        ("Pendências para validação humana", "pending_validations"),
    ):
        document.add_heading(heading, level=1)
        _add_docx_list(document, plan.get(key))

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_styles() -> dict[str, ParagraphStyle]:
    sample = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "PlanTitle",
            parent=sample["Title"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            alignment=TA_CENTER,
            textColor="#000000",
            spaceAfter=18,
        ),
        "h1": ParagraphStyle(
            "PlanHeading1",
            parent=sample["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=17,
            textColor="#000000",
            spaceBefore=12,
            spaceAfter=7,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "PlanHeading2",
            parent=sample["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11.5,
            leading=14,
            textColor="#000000",
            spaceBefore=9,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "PlanBody",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=15,
            textColor="#000000",
            spaceAfter=6,
        ),
    }


def _pdf_paragraph(label: str, value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(f"<b>{escape(label)}:</b> {escape(_text(value))}", style)


def _pdf_list(values: object, style: ParagraphStyle) -> ListFlowable:
    return ListFlowable(
        [ListItem(Paragraph(escape(value), style)) for value in _items(values)],
        bulletType="bullet",
        leftIndent=18,
        bulletFontName="Helvetica",
        bulletFontSize=8,
        spaceAfter=6,
    )


def _pdf_heading_list(
    heading: str,
    values: object,
    heading_style: ParagraphStyle,
    body_style: ParagraphStyle,
) -> KeepTogether:
    return KeepTogether(
        [
            Paragraph(escape(heading), heading_style),
            _pdf_list(values, body_style),
        ]
    )


def _page_number(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor("#555555")
    canvas.drawCentredString(letter[0] / 2, 0.42 * inch, str(document.page))
    canvas.restoreState()


def build_teaching_plan_pdf(plan: dict) -> bytes:
    output = BytesIO()
    styles = _pdf_styles()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.8 * inch,
        leftMargin=0.8 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.65 * inch,
        title=f"Plano de Ensino {_text(plan.get('title'))}",
        author="Assistente de Inteligência Curricular",
    )
    story = [
        Paragraph(
            escape(f"Plano de Ensino {_text(plan.get('title'))}"),
            styles["title"],
        ),
        Paragraph("Identificação e justificativa", styles["h1"]),
        _pdf_paragraph("Público alvo", plan.get("target_audience"), styles["body"]),
        _pdf_paragraph("Modalidade", plan.get("modality"), styles["body"]),
        _pdf_paragraph("Carga horária total", plan.get("total_workload"), styles["body"]),
        _pdf_paragraph("Justificativa", plan.get("rationale"), styles["body"]),
        _pdf_paragraph("Desempenho esperado", plan.get("expected_performance"), styles["body"]),
        Paragraph("Objetivos", styles["h1"]),
        _pdf_paragraph("Objetivo geral", plan.get("general_objective"), styles["body"]),
        _pdf_heading_list(
            "Objetivos específicos",
            plan.get("specific_objectives"),
            styles["h2"],
            styles["body"],
        ),
        Paragraph("Trilha de aprendizagem", styles["h1"]),
    ]

    steps = plan.get("learning_path") if isinstance(plan.get("learning_path"), list) else []
    if not steps:
        story.append(Paragraph("A definir mediante validação", styles["body"]))
    for index, step in enumerate(steps, start=1):
        sequence = step.get("sequence", index)
        story.extend(
            [
                Paragraph(
                    escape(f"Etapa {sequence} {_text(step.get('title'))}"),
                    styles["h2"],
                ),
                _pdf_paragraph("Objetivo", step.get("learning_objective"), styles["body"]),
            ]
        )
        for label, key in (
            ("Conteúdos", "contents"),
            ("Metodologia", "methodology"),
            ("Avaliação", "assessment"),
            ("Alinhamento com a MCN", "mcn_alignment"),
        ):
            story.append(
                _pdf_heading_list(
                    label,
                    step.get(key),
                    styles["h2"],
                    styles["body"],
                )
            )
        story.append(_pdf_paragraph("Carga horária", step.get("workload"), styles["body"]))

    for heading, key in (
        ("Estratégias de ensino", "teaching_methods"),
        ("Recursos necessários", "resources"),
        ("Avaliação da aprendizagem", "learning_assessment"),
        ("Avaliação de transferência e impacto", "transfer_impact_assessment"),
        ("Critérios de certificação", "certification_criteria"),
        ("Referências", "references"),
        ("Pendências para validação humana", "pending_validations"),
    ):
        story.append(
            _pdf_heading_list(
                heading,
                plan.get(key),
                styles["h1"],
                styles["body"],
            )
        )

    document.build(story, onFirstPage=_page_number, onLaterPages=_page_number)
    return output.getvalue()

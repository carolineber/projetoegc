from __future__ import annotations

import json
import sqlite3
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import load_workbook

XLSX_PATH = Path("espens_clean.xlsx")
DOCX_PATH = Path("_2.1._MCN 2026 SPB-3_07.04.2026.docx")
SQLITE_PATH = Path("data/local.db")
TABLE_NAME = "rag_docs"
WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _read_rows(xlsx_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        all_rows = list(sheet.iter_rows(values_only=True))
        if not all_rows:
            return []

        header_row = all_rows[0]
        headers = [
            str(value).strip() if value is not None else "" for value in header_row
        ]
        headers = [
            header if header else f"coluna_{idx + 1}"
            for idx, header in enumerate(headers)
        ]

        parsed_rows: list[dict[str, str]] = []
        for row in all_rows[1:]:
            row_dict: dict[str, str] = {}
            for idx, cell_value in enumerate(row):
                column_name = headers[idx]
                if not column_name:
                    continue
                row_dict[column_name] = (
                    "" if cell_value is None else str(cell_value).strip()
                )
            if any(row_dict.values()):
                parsed_rows.append(row_dict)
        return parsed_rows
    finally:
        workbook.close()


def _pick_column(columns: list[str], candidates: list[str], default: str) -> str:
    lowered = {col.lower(): col for col in columns}
    for cand in candidates:
        if cand in lowered:
            return lowered[cand]
    for col in columns:
        col_low = col.lower()
        if any(cand in col_low for cand in candidates):
            return col
    return default


def _cell_text(cell: ET.Element) -> str:
    paragraphs: list[str] = []
    for paragraph in cell.findall(f".//{WORD_NS}p"):
        text_value = "".join(
            node.text or "" for node in paragraph.iter(f"{WORD_NS}t")
        ).strip()
        if text_value:
            paragraphs.append(text_value)
    return "\n".join(paragraphs)


def _read_mcn_rows(docx_path: Path) -> list[dict[str, str]]:
    with zipfile.ZipFile(docx_path) as archive:
        document_root = ET.fromstring(archive.read("word/document.xml"))

    tables = document_root.findall(f".//{WORD_NS}tbl")
    if not tables:
        raise ValueError("Documento MCN sem tabela de dados.")

    rows = tables[0].findall(f"{WORD_NS}tr")
    if len(rows) < 3:
        raise ValueError("Tabela do documento MCN sem linhas de dados.")

    header_cells = rows[1].findall(f"{WORD_NS}tc")
    headers = [_cell_text(cell).replace("\n", " ").strip() for cell in header_cells]
    if headers and not headers[0]:
        headers[0] = "Competências (capacidades de/para)"

    parsed_rows: list[dict[str, str]] = []
    for row in rows[2:]:
        values = [_cell_text(cell) for cell in row.findall(f"{WORD_NS}tc")]
        if not any(values):
            continue

        parsed_row = {
            header: values[index].strip() if index < len(values) else ""
            for index, header in enumerate(headers)
            if header
        }
        if any(parsed_row.values()):
            parsed_rows.append(parsed_row)

    return parsed_rows


def _create_table(cursor: sqlite3.Cursor) -> None:
    cursor.execute(f"DROP TABLE IF EXISTS {TABLE_NAME}")
    cursor.execute(
        f"""
        CREATE TABLE {TABLE_NAME} (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_row INTEGER NOT NULL,
            source_name TEXT NOT NULL,
            source_type TEXT NOT NULL,
            titulo TEXT,
            cadastro TEXT,
            conteudo TEXT NOT NULL,
            metadata_json TEXT NOT NULL DEFAULT '{{}}'
        )
        """
    )


def _insert_record(
    cursor: sqlite3.Cursor,
    *,
    source_row: int,
    source_name: str,
    source_type: str,
    titulo: str,
    cadastro: str,
    conteudo: str,
    metadata: dict[str, str],
) -> None:
    cursor.execute(
        f"""
        INSERT INTO {TABLE_NAME}
            (
                source_row,
                source_name,
                source_type,
                titulo,
                cadastro,
                conteudo,
                metadata_json
            )
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            source_row,
            source_name,
            source_type,
            titulo,
            cadastro,
            conteudo,
            json.dumps(metadata, ensure_ascii=False),
        ),
    )


def main() -> None:
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Arquivo XLSX nao encontrado: {XLSX_PATH}")
    if not DOCX_PATH.exists():
        raise FileNotFoundError(f"Arquivo DOCX nao encontrado: {DOCX_PATH}")

    xlsx_rows = _read_rows(XLSX_PATH)
    if not xlsx_rows:
        raise ValueError("XLSX sem linhas de dados.")
    mcn_rows = _read_mcn_rows(DOCX_PATH)

    columns = list(xlsx_rows[0].keys())
    title_col = _pick_column(columns, ["tipo", "titulo"], columns[0])
    cadastro_col = _pick_column(columns, ["cadastro", "codigo"], columns[1])

    SQLITE_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(SQLITE_PATH)

    try:
        cursor = connection.cursor()
        _create_table(cursor)

        for idx, row in enumerate(xlsx_rows, start=1):
            titulo = row.get(title_col, "")
            cadastro = row.get(cadastro_col, "")
            conteudo = "\n".join(
                [
                    f"Fonte: {XLSX_PATH.name}",
                    *[f"{col}: {value}" for col, value in row.items() if value],
                ]
            )
            _insert_record(
                cursor,
                source_row=idx,
                source_name=XLSX_PATH.name,
                source_type="ações educativas ESPEN",
                titulo=titulo,
                cadastro=cadastro,
                conteudo=conteudo,
                metadata=row,
            )

        for idx, row in enumerate(mcn_rows, start=1):
            unit = row.get("Unidade Temática", "")
            critical_knowledge = row.get("Conhecimento Crítico e para Prática", "")
            titulo = " — ".join(value for value in (unit, critical_knowledge) if value)
            if not titulo:
                titulo = f"MCN 2026 — registro {idx}"
            cadastro = f"MCN_2026_{idx:04d}"
            conteudo = "\n".join(
                [
                    f"Fonte: {DOCX_PATH.name}",
                    *[f"{column}: {value}" for column, value in row.items() if value],
                ]
            )
            _insert_record(
                cursor,
                source_row=idx,
                source_name=DOCX_PATH.name,
                source_type="Matriz Curricular Nacional 2026",
                titulo=titulo,
                cadastro=cadastro,
                conteudo=conteudo,
                metadata=row,
            )

        connection.commit()
        print(
            f"SQLite criado em {SQLITE_PATH} com "
            f"{len(xlsx_rows)} registros de {XLSX_PATH.name} e "
            f"{len(mcn_rows)} registros de {DOCX_PATH.name}."
        )
    finally:
        connection.close()


if __name__ == "__main__":
    main()
